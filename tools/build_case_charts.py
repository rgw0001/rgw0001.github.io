#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Rebuild slovak-cases/ and czech-cases/ from the two case-chart workbooks.

Each worksheet becomes one tab of a self-contained HTML page, rendered as a
<table> that mirrors the workbook cell for cell: same merges, same fills, same
fonts, same borders, and Excel's own rule that text spills over its neighbours
only while they are empty.  Everything is measured in the workbook's units and
scaled once, so every tab lands inside one fixed-size box that needs no
scrolling.

    pip install openpyxl
    python3 tools/build_case_charts.py

Edit a workbook, re-run this, and the pages follow.  The card thumbnails on the
home page are separate; regenerate those by screenshotting each page and
cropping 4:3 from the tab bar down.
"""
import sys, html, colorsys, unicodedata
import openpyxl
from openpyxl.utils import get_column_letter

# clrScheme is stored dk1,lt1,dk2,lt2,accent1..6 but cells index it 0=lt1,1=dk1,...
THEME = ['FFFFFF', '000000', 'E7E6E6', '44546A',
         '5B9BD5', 'ED7D31', 'A5A5A5', 'FFC000', '4472C4', '70AD47']

def tinted(hexrgb, tint):
    """Excel applies a tint to the HSL luminance, not to the RGB channels."""
    r, g, b = (int(hexrgb[i:i+2], 16) / 255 for i in (0, 2, 4))
    h, l, s = colorsys.rgb_to_hls(r, g, b)
    l = l * (1 + tint) if tint < 0 else l * (1 - tint) + tint
    r, g, b = colorsys.hls_to_rgb(h, max(0.0, min(1.0, l)), s)
    return '#%02X%02X%02X' % (round(r*255), round(g*255), round(b*255))

def css_color(c):
    if c is None:
        return None
    t = getattr(c, 'type', None)
    if t == 'rgb':
        v = c.rgb
        if not isinstance(v, str) or len(v) != 8 or v[:2] == '00':
            return None                      # 00000000 is "no colour", not black
        return tinted(v[2:], c.tint or 0)
    if t == 'theme':
        try:
            return tinted(THEME[c.theme], c.tint or 0)
        except (IndexError, TypeError):
            return None
    if t == 'indexed':
        from openpyxl.styles.colors import COLOR_INDEX
        try:
            v = COLOR_INDEX[c.indexed]
        except (IndexError, TypeError):
            return None
        return tinted(v[2:], c.tint or 0) if v[:2] != '00' else None
    return None

def luminance(hexcolor):
    r, g, b = (int(hexcolor[i:i+2], 16) for i in (1, 3, 5))
    return 0.299 * r + 0.587 * g + 0.114 * b

BORDER_PX = {'hair': 1, 'thin': 1, 'dotted': 1, 'dashed': 1, 'double': 3,
             'medium': 2, 'mediumDashed': 2, 'thick': 3}
BORDER_STYLE = {'dotted': 'dotted', 'dashed': 'dashed', 'double': 'double'}

def border_css(side):
    if not side or not side.style:
        return None
    px = BORDER_PX.get(side.style, 1)
    st = BORDER_STYLE.get(side.style, 'solid')
    return '%dpx %s %s' % (px, st, css_color(side.color) or '#000')

def col_px(w):                 # Calibri 11 metrics, the usual approximation
    return round(w * 7) + 5

def col_widths(ws, ncols):
    """Excel stores <col min="3" max="7" width="13.7"/> as one run; openpyxl files
    that run under column C alone, so every run has to be expanded by hand."""
    default = ws.sheet_format.defaultColWidth or 8.43
    w = [default] * (ncols + 1)
    for dim in ws.column_dimensions.values():
        if not dim.width:
            continue
        lo, hi = dim.min or 1, dim.max or ncols
        for c in range(lo, min(hi, ncols) + 1):
            w[c] = dim.width
    return w

def has_text(ws, r, c):
    if c < 1 or c > ws.max_column:
        return False
    v = ws.cell(row=r, column=c).value
    return v is not None and str(v).strip() != ''

def sheet_table(ws, scale, ncols, nrows):
    merged = {}                       # anchor -> (rowspan, colspan)
    covered = set()                   # cells swallowed by an anchor
    for rng in ws.merged_cells.ranges:
        merged[(rng.min_row, rng.min_col)] = (rng.max_row - rng.min_row + 1,
                                              rng.max_col - rng.min_col + 1)
        for r in range(rng.min_row, rng.max_row + 1):
            for c in range(rng.min_col, rng.max_col + 1):
                if (r, c) != (rng.min_row, rng.min_col):
                    covered.add((r, c))

    widths = [col_px(w) * scale for w in col_widths(ws, ncols)[1:]]
    out = ['<table class="sheet" style="width:%.1fpx"><colgroup>' % sum(widths)]
    for w in widths:
        out.append('<col style="width:%.1fpx">' % w)
    out.append('</colgroup><tbody>')

    default_h = ws.sheet_format.defaultRowHeight or 15.0
    for r in range(1, nrows + 1):
        dim = ws.row_dimensions.get(r)
        h = dim.height if dim and dim.height else default_h
        out.append('<tr style="height:%.1fpx">' % (h * 4 / 3 * scale))
        for c in range(1, ncols + 1):
            if (r, c) in covered:
                continue
            cell = ws.cell(row=r, column=c)
            s = []
            bg = fg = None
            fill = cell.fill
            if fill is not None and fill.patternType == 'solid':
                bg = css_color(fill.fgColor)
                if bg:
                    s.append('background:' + bg)
            f = cell.font
            if f is not None:
                fg = css_color(f.color)
                if fg:
                    s.append('color:' + fg)
                if f.bold:
                    s.append('font-weight:700')
                if f.italic:
                    s.append('font-style:italic')
                if f.underline:
                    s.append('text-decoration:underline')
                if f.size:
                    s.append('font-size:%.1fpx' % (f.size * 4 / 3 * scale))
            al = cell.alignment
            if al is not None:
                if al.horizontal in ('center', 'right', 'left', 'justify'):
                    s.append('text-align:' + al.horizontal)
                if al.horizontal == 'right':
                    s.append('direction:rtl')
                if al.vertical in ('top', 'bottom', 'middle'):
                    s.append('vertical-align:' + al.vertical)
                if al.wrap_text:
                    s.append('white-space:normal')
                if al.indent:
                    s.append('padding-left:%.1fpx' % (al.indent * 8 * scale))
            b = cell.border
            if b is not None:
                for name, prop in (('left', 'border-left'), ('right', 'border-right'),
                                   ('top', 'border-top'), ('bottom', 'border-bottom')):
                    v = border_css(getattr(b, name))
                    if v:
                        s.append('%s:%s' % (prop, v))
            attrs = ''
            span = merged.get((r, c))
            if span:
                if span[0] > 1:
                    attrs += ' rowspan="%d"' % span[0]
                if span[1] > 1:
                    attrs += ' colspan="%d"' % span[1]
            if s:
                attrs += ' style="%s"' % ';'.join(s)
            v = cell.value
            txt = '' if v is None else html.escape(str(v)).replace('  ', ' &nbsp;')
            if txt:
                # Excel lets a cell's text spill over its neighbours only while
                # they are empty, and spills leftwards when the cell is
                # right-aligned.  Anything else gets clipped to the column.
                span_c = span[1] if span else 1
                ha = al.horizontal if al else None
                if ha == 'right':
                    blocked = has_text(ws, r, c - 1)
                elif ha == 'center':
                    blocked = has_text(ws, r, c - 1) or has_text(ws, r, c + span_c)
                else:
                    blocked = has_text(ws, r, c + span_c)
                if blocked:
                    txt = '<div class="clip">%s</div>' % txt
                elif bg and fg and luminance(fg) > 190 and luminance(bg) < 245:
                    # Pale text on a dark band would vanish the moment it spilled
                    # past the end of that band, so it carries the band with it.
                    txt = '<span class="spill" style="background:%s">%s</span>' % (bg, txt)
            out.append('<td%s>%s</td>' % (attrs, txt))
        out.append('</tr>')
    out.append('</tbody></table>')
    return ''.join(out)

def used_extent(ws):
    """Last row/column that carries either a value or a visible fill."""
    nr = nc = 0
    for row in ws.iter_rows():
        for cell in row:
            has_fill = cell.fill is not None and cell.fill.patternType == 'solid' \
                       and css_color(cell.fill.fgColor) is not None
            if cell.value is not None or has_fill:
                nr = max(nr, cell.row)
                nc = max(nc, cell.column)
    for rng in ws.merged_cells.ranges:
        nr, nc = max(nr, rng.max_row), max(nc, rng.max_col)
    return nr, nc

def natural_size(ws, nrows, ncols):
    default_h = ws.sheet_format.defaultRowHeight or 15.0
    w = sum(col_px(x) for x in col_widths(ws, ncols)[1:])
    h = sum((ws.row_dimensions.get(r).height
             if ws.row_dimensions.get(r) and ws.row_dimensions.get(r).height else default_h)
            for r in range(1, nrows + 1)) * 4 / 3
    return w, h

MAX_COL = 14          # column N; anything past it is emitted as a footnote
# The box is one fixed size for the whole page.  Its size is derived from the
# largest sheet in the workbook and one shared scale, so switching tabs never
# resizes anything and no chart needs scrolling inside the box.
FIT_W   = 1118        # most the box may ever be wide, in CSS px
FIT_H   = 898         # ...and tall

PAGE = u'''<!DOCTYPE html>
<html lang="{lang}">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>{title}</title>
<meta name="description" content="{desc}">
<link rel="preconnect" href="https://fonts.googleapis.com">
<link rel="preconnect" href="https://fonts.gstatic.com" crossorigin>
<!-- Carlito is metric-compatible with Calibri, so the grid measures out exactly
     as it does in the workbook even on machines with no Office fonts. -->
<link href="https://fonts.googleapis.com/css2?family=Carlito:ital,wght@0,400;0,700;1,400&family=IBM+Plex+Mono:wght@400;500&family=IBM+Plex+Serif:ital,wght@0,400;0,500;1,400&display=swap" rel="stylesheet">
<style>
  :root{{
    --paper:#F1F3F0; --ink:#14181A; --muted:#5E6763; --rule:#D2D8D2;
    --accent:#146A66;
    --serif:"IBM Plex Serif",Georgia,serif;
    --mono:"IBM Plex Mono",ui-monospace,monospace;
    /* The chart itself keeps Excel's own face, so it reads like the workbook. */
    --sheet:Carlito,Calibri,"Segoe UI",system-ui,-apple-system,"Helvetica Neue",Arial,sans-serif;
    --box-w:{box_w}px; --box-h:{box_h}px;
  }}
  *{{box-sizing:border-box}}
  body{{margin:0;background:var(--paper);color:var(--ink);
       font-family:var(--serif);font-size:17px;line-height:1.6}}

  /* The page is only ever a frame around the box, so the wrap tracks the box. */
  .wrap{{max-width:var(--box-w);margin:0 auto;padding:34px 0 64px}}
  @media (max-width:1010px){{ .wrap{{max-width:100%;padding:24px 16px 48px}} }}

  h1{{font-family:var(--serif);font-weight:400;font-size:1.6rem;margin:0 0 6px}}
  .sub{{color:var(--muted);font-size:15px;margin:0 0 10px;max-width:56em}}
  /* The lead sentence stays a paragraph; the points after it become a list. */
  .sub-list{{color:var(--muted);font-size:15px;line-height:1.6;
            margin:0 0 22px;max-width:56em;padding-left:1.15em}}
  .sub-list li{{margin:0 0 5px}}
  .sub-list li:last-child{{margin-bottom:0}}
  .back{{font-family:var(--mono);font-size:12px;letter-spacing:.08em;text-transform:uppercase;
        color:var(--muted);text-decoration:none;display:inline-block;margin:0 0 20px}}
  .back:hover,.back:focus-visible{{color:var(--accent)}}

  /* ---------- tabs ---------- */
  .tabs{{display:flex;flex-wrap:wrap;gap:6px;margin:0;padding:14px 0 12px;list-style:none;
         position:sticky;top:0;z-index:5;background:var(--paper)}}
  .tabs button{{
    font-family:var(--mono);font-size:11.5px;letter-spacing:.06em;text-transform:uppercase;
    font-weight:500;padding:7px 13px;border:1px solid #C6CEC8;background:#fff;
    color:var(--ink);border-radius:999px;cursor:pointer;transition:background .15s,border-color .15s,color .15s}}
  .tabs button:hover{{border-color:var(--accent);color:var(--accent)}}
  .tabs button[aria-selected="true"]{{background:var(--accent);border-color:var(--accent);color:#fff}}

  /* ---------- the fixed box ---------- */
  .box{{
    width:var(--box-w);height:var(--box-h);max-width:100%;
    border:1px solid var(--rule);border-radius:10px;background:#fff;
    box-shadow:0 6px 22px rgba(20,24,26,.07);
    overflow:auto;-webkit-overflow-scrolling:touch}}
  .panel[hidden]{{display:none}}
  .note{{font-family:var(--sheet);font-size:11px;line-height:1.3;color:#5E6763;
         margin:0;padding:5px 8px 4px}}

  /* ---------- the workbook grid ---------- */
  table.sheet{{border-collapse:collapse;table-layout:fixed;font-family:var(--sheet);
               font-size:{base_px:.1f}px;line-height:1.05;color:#000}}
  table.sheet td{{padding:0 {pad:.1f}px;vertical-align:bottom;white-space:nowrap;overflow:visible}}
  /* a cell whose neighbour is occupied cannot spill, so it clips instead */
  table.sheet .clip{{overflow:hidden}}
  /* the span carries the cell's own fill; inline-block keeps that fill from
     bleeding out of the short rows the workbook uses as gutters */
  table.sheet .spill{{display:inline-block;line-height:1}}

  .hint{{font-family:var(--mono);font-size:11.5px;color:var(--muted);margin:12px 0 0}}
  :focus-visible{{outline:2px solid var(--accent);outline-offset:3px}}
</style>
</head>
<body>
<div class="wrap">

  <a class="back" href="../">&larr; Robert G. West</a>
  <h1>{h1}</h1>
  {sub}

  <div class="tabs" role="tablist" aria-label="{tablist_label}">
{tabs}
  </div>

  <div class="box">
{panels}
  </div>

  <p class="hint">Arrow keys move between tabs, and each tab keeps its own link.</p>
</div>

<script>
// One tab per worksheet.  The chosen tab is mirrored into the URL hash so an
// individual case can be linked to directly.
(function(){{
  var tabs   = Array.prototype.slice.call(document.querySelectorAll('.tabs button'));
  var panels = Array.prototype.slice.call(document.querySelectorAll('.panel'));
  var box    = document.querySelector('.box');

  function show(i, focus){{
    tabs.forEach(function(t, j){{
      t.setAttribute('aria-selected', j === i ? 'true' : 'false');
      t.tabIndex = j === i ? 0 : -1;
      panels[j].hidden = j !== i;
    }});
    box.scrollTop = 0; box.scrollLeft = 0;
    if(focus) tabs[i].focus();
    history.replaceState(null, '', '#' + tabs[i].dataset.key);
  }}

  tabs.forEach(function(t, i){{
    t.addEventListener('click', function(){{ show(i, false); }});
    t.addEventListener('keydown', function(e){{
      var d = e.key === 'ArrowRight' ? 1 : e.key === 'ArrowLeft' ? -1 : 0;
      if(!d) return;
      e.preventDefault();
      show((i + d + tabs.length) % tabs.length, true);
    }});
  }});

  var start = tabs.findIndex(function(t){{ return '#' + t.dataset.key === location.hash; }});
  show(start < 0 ? 0 : start, false);
}})();
</script>

</body>
</html>
'''

import unicodedata

def slug(s):
    """ASCII-fold, so a case links as #instrumental rather than #in%C5%A1trument%C3%A1l."""
    s = unicodedata.normalize('NFKD', s.lower())
    return ''.join(ch if (ch.isascii() and ch.isalnum()) else '-'
                   for ch in s if not unicodedata.combining(ch)).strip('-')

def sub_block(lead, *points):
    """Lead sentence as a paragraph, then one <li> per point."""
    items = '\n'.join('    <li>%s</li>' % p for p in points)
    return '<p class="sub">%s</p>\n\n  <ul class="sub-list">\n%s\n  </ul>' % (lead, items)


def build(xlsx, out, lang, title, desc, h1, sub, tablist_label):
    wb = openpyxl.load_workbook(xlsx)

    # One scale for every sheet, so the grid never changes size between tabs.
    sizes = [natural_size(ws, used_extent(ws)[0], MAX_COL) for ws in wb.worksheets]
    widest, tallest = max(w for w, _ in sizes), max(h for _, h in sizes)
    scale = min(FIT_W / widest, FIT_H / tallest)
    # +6 absorbs sub-pixel rounding, so no tab ever gains a scrollbar
    box_w, box_h = round(widest * scale) + 6, round(tallest * scale) + 6

    tabs, panels, note_h = [], [], 0
    for i, ws in enumerate(wb.worksheets):
        nrows, ncols = used_extent(ws)
        ncols = min(ncols, MAX_COL)
        key = slug(ws.title)
        label = ws.title if ws.title.istitle() else ws.title.capitalize()
        tabs.append('    <button type="button" role="tab" data-key="%s" '
                    'aria-controls="p-%s" aria-selected="false" tabindex="-1">%s</button>'
                    % (key, key, html.escape(label)))

        strays = [str(c.value) for row in ws.iter_rows(min_col=MAX_COL + 1)
                  for c in row if c.value is not None]
        note = ''
        if strays:
            note = '<p class="note">%s</p>' % html.escape(' · '.join(strays))
            note_h = 30          # keeps a sheet with a footnote inside the box

        panels.append('    <div class="panel" id="p-%s" role="tabpanel" hidden>%s%s</div>'
                      % (key, sheet_table(ws, scale, ncols, nrows), note))

    box_h += note_h
    page = PAGE.format(lang=lang, title=html.escape(title), desc=html.escape(desc),
                       h1=html.escape(h1), sub=sub, tablist_label=html.escape(tablist_label),
                       box_w=box_w, box_h=box_h, base_px=11 * 4 / 3 * scale, pad=2 * scale,
                       tabs='\n'.join(tabs), panels='\n'.join(panels))
    with open(out, 'w', encoding='utf-8') as fh:
        fh.write(page)
    print('wrote', out, '(scale %.3f)' % scale)

if __name__ == '__main__':
    root = '/Users/rgw0001/Git/rgw0001/rgw0001.github.io/'
    build('/Users/rgw0001/Documents/Slovenské Pady.xlsx', root + 'slovak-cases/index.html',
          'sk', 'Slovenské pády — Slovak case chart',
          'The six Slovak cases: nouns, pronouns, adjectives and possessives, one tab per case.',
          'Slovenské pády',
          sub_block(
          'Select a tab to show the summary of the comparison of nominative form of '
          'singular and plural words with the case, indicated by the red letters at the '
          'ends of the words.',
          'The letters written at the top of the columns indicate that words ending in '
          'these letters follow the form in the rows below them for the male gender.',
          'The two words on the right indicate the question that is asked for the case: '
          'who and what? This is how a case is identified.',
          'The prepositions and conditions for using this case are then written in the '
          'black box. Truly, there are four genders for us English speakers: male living '
          '(includes the dead fish on your plate and Lego minifigures), male non-living, '
          'female, and neutral.',
          'Pronouns and reflexive pronouns change too.',
          'At the bottom are adjectives, which also change for each case and gender.'),
          'Slovak cases')
    build('/Users/rgw0001/Documents/České Pady.xlsx', root + 'czech-cases/index.html',
          'cs', 'České pády — Czech case chart',
          'The seven Czech cases: nouns, pronouns, adjectives and possessives, one tab per case, '
          'plus a blank practice sheet.',
          'České pády',
          sub_block(
          'Select a tab to show the summary of the comparison of nominative form of words '
          'with the case, indicated by the red letters at the ends of the words.',
          'The letters written at the top of the columns indicate that words ending in '
          'these letters follow the form in the rows below them for the male gender.',
          'The two words on the right indicate the question that is asked for the case: '
          'who and what? This is how a case is identified.',
          'The prepositions and conditions for using this case are then written in the '
          'black box. Truly, there are four genders for us English speakers: male living '
          '(includes the dead fish on your plate and Lego minifigures), male non-living, '
          'female, and neutral.',
          'Pronouns and reflexive pronouns change too.',
          'At the bottom are adjectives, which also change for each case and gender.'),
          'Czech cases')
